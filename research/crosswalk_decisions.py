#!/usr/bin/env python3
"""A returned approval sheet, read back as claim decisions (ADR 0041, zone 4). Each row of the sheet stands for the
claims it shows (docs/crosswalk/review/<step>.rows.json); a row with a decision becomes a decision on each of them,
by the signatory who returned it and by nobody else. A draft nobody decides stays a proposal.

python3 research/crosswalk_decisions.py MATH.NUMBER.8 <the sheet, downloaded as .csv or .xlsx> --by Akanksha
"""

import argparse
import datetime
import json
import sys
from pathlib import Path

from crosswalk_review import CW, R, load, slug

CHOICES = {"Approve": "approved", "Change": "revised", "Reject": "rejected"}


def decisions(step, path, by):
    """A returned sheet's decisions, as claim decisions by the signatory who returned it; refused otherwise."""
    t = load()
    if by not in {s["educator"] for s in t["signatory"]}:
        raise SystemExit(f"{by} is not a signatory; only a signatory's decisions count")
    ids = json.loads((CW / "review" / f"{slug(step)}.rows.json").read_text())
    if str(path).endswith(".csv"):
        import csv

        with open(path, newline="") as f:
            grids = [(Path(path).stem, list(csv.reader(f)))]
    else:
        from openpyxl import load_workbook

        grids = [
            (ws.title, [list(r) for r in ws.iter_rows(values_only=True)])
            for ws in load_workbook(path).worksheets
        ]
    return read_decisions(grids, ids, by, Path(path).name)


def read_decisions(grids, ids, by, source, now=None):
    """The rows with a decision, each turned into a decision on every claim the row stands for. `grids` is a list of
    (tab, rows); a tab's header is the first row that has a Decision column."""
    now = now or datetime.datetime.now(datetime.UTC).isoformat(timespec="seconds")
    out, faults = [], []
    for title, grid in grids:
        top = next((i for i, r in enumerate(grid) if "Decision" in r), None)
        if top is None:
            continue
        header = grid[top]
        at, note = header.index("Decision"), header.index("Comment")
        for cells in grid[top + 1 :]:
            cells = list(cells) + [None] * (len(header) - len(cells))
            row, choice = cells[0], (cells[at] or "").strip()
            if not row or not choice:
                continue
            if choice not in CHOICES or row not in ids:
                faults.append(
                    f"{title} {row}: '{choice}' is not Approve, Change or Reject, or the row is unknown"
                )
                continue
            comment = (cells[note] or "").strip()
            if CHOICES[choice] != "approved" and not comment:
                comment = "(no comment given)"
            out += [
                {
                    "claim_id": c,
                    "decision": CHOICES[choice],
                    "decided_by": by,
                    "decided_at": now,
                    "note": comment,
                    "source": f"{source}, {title} row {row}",
                }
                for c in ids[row]
            ]
    return out, faults


def main(argv):
    ap = argparse.ArgumentParser()
    ap.add_argument("step")
    ap.add_argument("file")
    ap.add_argument("--by", required=True)
    a = ap.parse_args(argv)
    got, faults = decisions(a.step, a.file, a.by)
    for f in faults:
        print("FAULT:", f)
    if faults:
        return 1
    out = CW / "decisions" / f"{slug(a.step)}-{datetime.date.today().isoformat()}.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(got, ensure_ascii=False, indent=1) + "\n")
    print(
        f"{len(got)} claim decisions by {a.by} -> {out.relative_to(R)}; research/crosswalk_build.py applies them"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
